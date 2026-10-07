from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from sqlalchemy.orm import Session

from backend.app.ai.gemini import gemini_service
from backend.app.ai.resume_parser import resume_parser
from backend.app.ai.taxonomy import taxonomy
from backend.app.config import get_settings
from backend.app.database.session import get_db
from backend.app.models.entities import Resume, Skill, User, UserSkill
from backend.app.schemas.schemas import ResumeAnalysisResponse, ResumeParseOutput
from backend.app.security import get_current_user
from backend.app.storage import resume_storage

router = APIRouter(prefix="/api/resume", tags=["Resume"])


@router.post("/upload", response_model=ResumeAnalysisResponse)
async def upload_resume(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    filename = file.filename or "resume"
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix not in {"pdf", "docx", "txt"}:
        raise HTTPException(status_code=415, detail="Upload a PDF, DOCX, or TXT resume")
    content = await file.read(get_settings().resume_max_bytes + 1)
    if len(content) > get_settings().resume_max_bytes:
        raise HTTPException(status_code=413, detail="Resume exceeds the configured upload limit")
    if suffix == "pdf":
        raw_text = resume_parser.extract_text_from_pdf(content)
    elif suffix == "docx":
        try:
            from docx import Document
            from io import BytesIO
            from zipfile import ZipFile
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 2048 or sum(entry.file_size for entry in entries) > 50 * 1024 * 1024:
                    raise HTTPException(status_code=413, detail="DOCX expands beyond the allowed document size")
                if "word/document.xml" not in archive.namelist():
                    raise HTTPException(status_code=422, detail="Invalid DOCX document")
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            paragraphs.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells if cell.text.strip())
            raw_text = "\n".join(paragraphs)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=422, detail="Could not read the DOCX resume") from exc
    else:
        raw_text = content.decode("utf-8", errors="replace")
    if not raw_text.strip():
        raise HTTPException(status_code=422, detail="No readable text was found in this resume")

    parsed_data = gemini_service.parse_resume(raw_text) or resume_parser.parse_resume_content(raw_text)
    parsed = ResumeParseOutput.model_validate(parsed_data)
    try:
        object_path = resume_storage.put(user.id, filename, content)
    except Exception as exc:
        raise HTTPException(status_code=503, detail="Resume storage is unavailable") from exc

    resume_record = Resume(
        user_id=user.id,
        filename=filename,
        file_path=object_path,
        raw_text=raw_text,
        parsed_data=parsed.model_dump(),
    )
    db.add(resume_record)
    db.flush()
    extracted_canonical = []
    for skill_info in parsed.skills:
        canonical = taxonomy.normalize(skill_info.name) or skill_info.name.strip()
        if not canonical:
            continue
        extracted_canonical.append(canonical)
        skill = db.query(Skill).filter(Skill.name.ilike(canonical)).first()
        if not skill:
            skill = Skill(
                name=canonical,
                category=skill_info.category or "Technical",
                importance=skill_info.importance or "Medium",
            )
            db.add(skill)
            db.flush()
        existing = db.query(UserSkill).filter(UserSkill.user_id == user.id, UserSkill.skill_id == skill.id).first()
        if not existing:
            db.add(UserSkill(
                user_id=user.id,
                skill_id=skill.id,
                proficiency=skill_info.proficiency or "Moderate",
                verified=False,
                source="Resume extraction",
            ))
    db.commit()
    db.refresh(resume_record)
    return ResumeAnalysisResponse(
        resume_id=resume_record.id,
        filename=resume_record.filename,
        parsed_data=parsed,
        extracted_skills_count=len(extracted_canonical),
        matched_canonical_skills=extracted_canonical,
    )


@router.get("/latest", response_model=ResumeAnalysisResponse)
def get_latest_resume(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    resume = (
        db.query(Resume)
        .filter(Resume.user_id == user.id)
        .order_by(Resume.uploaded_at.desc())
        .first()
    )
    if not resume:
        raise HTTPException(status_code=404, detail="No resume uploaded")
    parsed = ResumeParseOutput.model_validate(resume.parsed_data or {})
    canonical = [skill.name for skill in parsed.skills]
    return ResumeAnalysisResponse(
        resume_id=resume.id,
        filename=resume.filename,
        parsed_data=parsed,
        extracted_skills_count=len(canonical),
        matched_canonical_skills=canonical,
    )
