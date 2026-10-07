from backend.app.models.entities import Resume, Roadmap, RoadmapItem, User, UserSession, UserSkill


def test_public_health_and_protected_routes(api_client):
    assert api_client.get("/api/health").json()["status"] == "healthy"
    assert api_client.get("/api/profile").status_code == 401
    assert api_client.get("/api/dashboard").status_code == 401
    assert api_client.get("/api/skills/user").status_code == 401
    assert api_client.post("/api/skills/user", json={"skill_name": "Python"}).status_code == 403


def test_production_configuration_fails_fast_when_integrations_are_missing():
    from types import SimpleNamespace
    import pytest
    from backend.main import validate_runtime_settings

    incomplete = SimpleNamespace(
        is_production=True,
        session_secret="short",
        database_url="sqlite:///careerbridge.db",
        s3_bucket="",
        google_client_id="",
        google_client_secret="",
        gemini_api_key="",
        adzuna_app_id="",
        adzuna_app_key="",
        session_cookie_secure=False,
        frontend_url="http://localhost:3000",
        google_redirect_uri="http://localhost:8000/api/auth/google/callback",
    )
    with pytest.raises(RuntimeError, match="Production configuration is incomplete") as error:
        validate_runtime_settings(incomplete)
    for setting in ("DATABASE_URL", "GOOGLE_CLIENT_ID", "GEMINI_API_KEY", "ADZUNA_APP_ID", "S3_BUCKET"):
        assert setting in str(error.value)
    ready = SimpleNamespace(**{
        **vars(incomplete),
        "session_secret": "s" * 48,
        "database_url": "postgresql+psycopg://app:secret@db/careerbridge",
        "s3_bucket": "careerbridge-private-resumes",
        "google_client_id": "google-client",
        "google_client_secret": "google-secret",
        "gemini_api_key": "gemini-key",
        "adzuna_app_id": "adzuna-id",
        "adzuna_app_key": "adzuna-key",
        "session_cookie_secure": True,
        "frontend_url": "https://careerbridge.example",
        "google_redirect_uri": "https://api.careerbridge.example/api/auth/google/callback",
    })
    assert validate_runtime_settings(ready) is None

    unsafe = SimpleNamespace(**{
        **vars(ready),
        "session_secret": "local-development-session-secret-change-before-deploy",
        "database_url": "postgresql+psycopg://app:local-development-only@db/careerbridge",
        "s3_endpoint_url": "http://minio:9000",
    })
    with pytest.raises(RuntimeError, match="Production configuration is incomplete") as error:
        validate_runtime_settings(unsafe)
    assert "SESSION_SECRET" in str(error.value)
    assert "POSTGRES_PASSWORD" in str(error.value)
    assert "S3_ENDPOINT_URL" in str(error.value)


def test_google_login_configuration_and_logout(api_client, account, monkeypatch):
    from backend.app.api import auth
    monkeypatch.setattr(auth.settings, "google_client_id", "")
    monkeypatch.setattr(auth.settings, "google_client_secret", "")
    assert api_client.get("/api/auth/google").status_code == 503
    assert api_client.get("/api/auth/me").status_code == 200
    assert api_client.post("/api/auth/logout").status_code == 200
    assert api_client.get("/api/auth/me").status_code == 401


def test_google_callback_creates_verified_account_and_session(api_client, test_db, monkeypatch):
    from backend.app.api import auth

    async def verified_google_user(request):
        return {"userinfo": {
            "sub": "google-new-user",
            "email": "new@example.test",
            "email_verified": True,
            "name": "New Candidate",
        }}

    monkeypatch.setattr(auth.oauth.google, "authorize_access_token", verified_google_user)
    response = api_client.get("/api/auth/google/callback")
    assert response.status_code == 302
    assert response.headers["location"].endswith("/auth/complete")
    assert response.headers["set-cookie"]
    identity = api_client.get("/api/auth/me").json()
    assert identity["email"] == "new@example.test"
    assert identity["full_name"] == "New Candidate"
    user = test_db.query(User).filter_by(email="new@example.test").one()
    assert user.google_subject == "google-new-user"
    assert user.profile is not None
    assert user.profile.current_title is None
    assert user.profile.target_role_id is None
    assert user.profile.target_role_title is None
    assert user.profile.experience_level is None
    assert user.profile.education_level is None


def test_identity_profile_and_preferences(api_client, account):
    assert api_client.get("/api/auth/me").json() == {
        "id": "user-a", "email": "a@example.test", "full_name": "A Candidate"
    }
    profile = api_client.get("/api/profile").json()
    assert profile["target_role_id"] == "role-test"
    updated = api_client.put("/api/profile", json={"current_title": "QA Engineer"})
    assert updated.status_code == 200
    assert updated.json()["current_title"] == "QA Engineer"

    prefs = api_client.put("/api/preferences", json={"salary_expectation": "₹8,00,000–₹12,00,000", "remote_only": True})
    assert prefs.status_code == 200
    assert api_client.get("/api/preferences").json()["remote_only"] is True


def test_unconfigured_profile_does_not_silently_choose_a_career(api_client, account, test_db):
    from backend.app.models.entities import Profile

    profile = test_db.query(Profile).filter_by(user_id="user-a").one()
    profile.target_role_id = None
    profile.target_role_title = None
    test_db.commit()

    dashboard = api_client.get("/api/dashboard")
    assert dashboard.status_code == 200
    assert dashboard.json()["target_role"] is None
    assert dashboard.json()["profile"]["target_role_title"] is None
    assert api_client.get("/api/skill-gap/current").status_code == 409
    assert api_client.get("/api/roadmap/current").status_code == 409


def test_readiness_and_public_catalogs(api_client, account):
    assert api_client.get("/api/ready").json() == {"status": "ready", "database": "reachable"}
    skills = api_client.get("/api/skills")
    careers = api_client.get("/api/careers")
    assert skills.status_code == 200
    assert careers.status_code == 200
    assert any(item["name"] == "Python" for item in skills.json())
    assert any(item["id"] == "role-test" for item in careers.json())


def test_skills_career_match_and_skill_gap(api_client, account):
    skills = api_client.get("/api/skills/user")
    assert skills.status_code == 200
    assert skills.json()[0]["skill_name"] == "Python"
    added = api_client.post("/api/skills/user", json={"skill_name": "SQL", "proficiency": "Familiar"})
    assert added.status_code == 200
    assert added.json()["proficiency"] == "Familiar"
    assert api_client.post("/api/skills/user", json={"skill_name": "SQL", "proficiency": "Novice"}).status_code == 422

    assert api_client.get("/api/careers/role-test").status_code == 200
    recs = api_client.post("/api/careers/recommend").json()["recommendations"]
    assert recs[0]["role_id"] == "role-test"
    assert recs[0]["match_score"] == 65.0
    analysis = api_client.post("/api/skill-gap/analyze", json={"role_id": "role-test"})
    assert analysis.status_code == 200
    assert analysis.json()["readiness_score"] == 65.0
    current_analysis = api_client.get("/api/skill-gap/current")
    assert current_analysis.status_code == 200
    assert current_analysis.json()["role_id"] == "role-test"

    skill_id = added.json()["id"]
    assert api_client.delete(f"/api/skills/user/{skill_id}").status_code == 200
    assert api_client.delete(f"/api/skills/user/{skill_id}").status_code == 404


def test_roadmap_progress_and_dashboard(api_client, account, monkeypatch):
    from backend.app.ai.gemini import gemini_service

    monkeypatch.setattr(gemini_service, "generate_roadmap", lambda *args: None)
    roadmap = api_client.get("/api/roadmap/current")
    assert roadmap.status_code == 200
    first = roadmap.json()["phases"][0]
    assert first["id"]
    assert first["learning_goals"]
    update = api_client.patch(
        f"/api/roadmap/items/{first['id']}",
        json={"completed_goals": [first["learning_goals"][0]]},
    )
    assert update.status_code == 200
    assert update.json()["current_progress"] > 0
    assert update.json()["phases"][0]["completed_goals"] == [first["learning_goals"][0]]
    assert api_client.patch(f"/api/roadmap/items/{first['id']}", json={"completed_goals": ["not a real goal"]}).status_code == 422

    dashboard = api_client.get("/api/dashboard")
    assert dashboard.status_code == 200
    assert dashboard.json()["readiness_score"] == 0.0
    assert dashboard.json()["roadmap_progress"] > 0


def test_roadmap_can_be_regenerated_for_selected_role(api_client, account, monkeypatch):
    from backend.app.ai.gemini import gemini_service

    monkeypatch.setattr(gemini_service, "generate_roadmap", lambda *args: None)
    response = api_client.post("/api/roadmap/generate?role_id=role-test")
    assert response.status_code == 200
    assert response.json()["role_id"] == "role-test"
    assert len(response.json()["phases"]) == 5
    assert api_client.get("/api/roadmap/current").json()["id"] == response.json()["id"]


def test_resume_upload_and_latest(api_client, account, monkeypatch):
    from backend.app.storage import resume_storage
    from backend.app.ai.gemini import gemini_service

    monkeypatch.setattr(resume_storage, "put", lambda *args: "local/test-resume.txt")
    monkeypatch.setattr(gemini_service, "parse_resume", lambda text: None)
    upload = api_client.post(
        "/api/resume/upload",
        files={"file": ("resume.txt", b"A Candidate\nSkills: Python, SQL\nBachelor of Science, 2023", "text/plain")},
    )
    assert upload.status_code == 200
    assert upload.json()["filename"] == "resume.txt"
    assert upload.json()["parsed_data"]["skills"]
    latest = api_client.get("/api/resume/latest")
    assert latest.status_code == 200
    assert latest.json()["resume_id"] == upload.json()["resume_id"]
    assert api_client.post("/api/resume/upload", files={"file": ("bad.exe", b"x", "application/octet-stream")}).status_code == 415


def test_docx_resume_upload_and_invalid_or_oversized_documents(api_client, account, monkeypatch):
    from io import BytesIO
    from types import SimpleNamespace
    from docx import Document
    from backend.app.api import resume
    from backend.app.ai.gemini import gemini_service
    from backend.app.storage import resume_storage

    document = Document()
    document.add_paragraph("A Candidate")
    document.add_paragraph("Skills: Python and SQL")
    content = BytesIO()
    document.save(content)
    monkeypatch.setattr(resume_storage, "put", lambda *args: "local/test-resume.docx")
    monkeypatch.setattr(gemini_service, "parse_resume", lambda text: None)

    uploaded = api_client.post(
        "/api/resume/upload",
        files={"file": ("resume.docx", content.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert uploaded.status_code == 200
    assert "Python" in uploaded.json()["matched_canonical_skills"]

    invalid_docx = api_client.post(
        "/api/resume/upload",
        files={"file": ("broken.docx", b"not a zip archive", "application/octet-stream")},
    )
    assert invalid_docx.status_code == 422
    invalid_pdf = api_client.post(
        "/api/resume/upload",
        files={"file": ("broken.pdf", b"not a PDF", "application/pdf")},
    )
    assert invalid_pdf.status_code == 422

    monkeypatch.setattr(resume, "get_settings", lambda: SimpleNamespace(resume_max_bytes=16))
    oversized = api_client.post(
        "/api/resume/upload",
        files={"file": ("large.txt", b"A Candidate\nSkills: Python", "text/plain")},
    )
    assert oversized.status_code == 413


def test_user_data_isolation(api_client, test_db, account):
    second = User(id="user-b", email="b@example.test", full_name="B Candidate", google_subject="google-b")
    test_db.add(second)
    test_db.flush()
    import hashlib, secrets
    token = secrets.token_urlsafe(48)
    from datetime import datetime, timedelta
    test_db.add(UserSession(user_id=second.id, token_hash=hashlib.sha256(token.encode()).hexdigest(), expires_at=datetime.utcnow() + timedelta(days=1)))
    test_db.commit()
    api_client.cookies.set("cb_session", token)
    result = api_client.get("/api/skills/user")
    assert result.status_code == 200
    assert result.json() == []


def test_csrf_and_expired_session_are_rejected(api_client, test_db, account):
    from datetime import datetime, timedelta

    api_client.cookies.delete("cb_csrf")
    assert api_client.post("/api/skills/user", json={"skill_name": "Go"}).status_code == 403
    session = test_db.query(UserSession).filter_by(user_id="user-a").one()
    session.expires_at = datetime.utcnow() - timedelta(minutes=1)
    test_db.commit()
    assert api_client.get("/api/profile").status_code == 401


def test_resume_and_roadmap_records_are_isolated(api_client, test_db, account):
    from datetime import datetime

    second = User(id="user-b", email="b@example.test", full_name="B Candidate", google_subject="google-b")
    test_db.add(second)
    test_db.flush()
    roadmap = Roadmap(user_id=second.id, role_id=account["role"].id, title="B roadmap")
    test_db.add(roadmap)
    test_db.flush()
    phase = RoadmapItem(
        roadmap_id=roadmap.id,
        phase_number=1,
        phase_name="Foundation",
        title="B phase",
        description="B-only learning phase",
        skills=["Python"],
        milestone_projects=[],
        learning_goals=["B-only goal"],
        completed_goals=[],
    )
    test_db.add(phase)
    test_db.add(Resume(
        user_id=second.id,
        filename="private.txt",
        file_path="private/object",
        parsed_data={"skills": [], "education": [], "experience": [], "projects": [], "certifications": []},
        uploaded_at=datetime.utcnow(),
    ))
    test_db.commit()

    assert api_client.get("/api/resume/latest").status_code == 404
    assert api_client.patch(
        f"/api/roadmap/items/{phase.id}", json={"completed_goals": ["B-only goal"]}
    ).status_code == 404
    assert test_db.query(RoadmapItem).filter_by(id=phase.id).one().completed_goals == []


def test_opportunity_api_uses_authenticated_match_results(api_client, account, monkeypatch):
    from backend.app.api import opportunities
    expected = [{
        "id": "job-1", "title": "Python Engineer", "company": "Example Co", "location": "Kolkata",
        "type": "Full-time", "salary_range": None, "experience_level": "Not specified",
        "description": "Python work", "apply_url": "https://jobs.example.test/1",
        "required_skills": ["Python"], "preferred_skills": [], "match_score": 85.0,
        "matched_skills": ["Python"], "missing_skills": [],
    }]
    monkeypatch.setattr(opportunities, "matched_opportunities", lambda *args, **kwargs: expected)
    response = api_client.get("/api/opportunities?type_filter=Full-time")
    assert response.status_code == 200
    assert response.json() == expected
